"""
条件选择（图层属性空间连接）服务测试

使用合成的 shapely 要素（不依赖真实文件）；run_attribute_join 通过 monkeypatch
替换 load_layer_features。导出目录重定向到 tmp_path。
"""
import pytest
from pathlib import Path

from shapely.geometry import Point, Polygon, LineString, box
from openpyxl import load_workbook

from app.core.config import settings
from app.services import attribute_join_service as svc
from app.services.attribute_join_service import (
    JoinFeature,
    run_attribute_join,
    build_rows_and_export,
    _spatial_match,
    valid_lonlat,
    RELATION_COL,
    RELATION_IN,
    RELATION_OUT,
    RELATION_DANGLING,
    RELATION_SKIP,
    REASON_MISSING_COORD,
)


def _square(cx: float, cy: float, half: float = 0.02) -> Polygon:
    return box(cx - half, cy - half, cx + half, cy + half)


@pytest.fixture(autouse=True)
def _export_dir(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(settings, "EXPORT_DIR", tmp_path)
    return tmp_path


def _patch_loader(monkeypatch, f1, f2, s1=None, s2=None):
    seq = iter([(f1, s1 or []), (f2, s2 or [])])
    monkeypatch.setattr(svc, "load_layer_features", lambda ref, g=None: next(seq))


def _read_sheet(path):
    ws = load_workbook(path).active
    rows = list(ws.iter_rows(values_only=True))
    return rows[0], rows[1:]


# ============================ 空间匹配 ============================

def test_all_points_inside():
    f1 = [JoinFeature("a", {"小区": "c1"}, Point(116.30, 39.90)),
          JoinFeature("b", {"小区": "c2"}, Point(116.31, 39.91))]
    f2 = [JoinFeature("p1", {"县名": "东城区"}, _square(116.305, 39.905, 0.03))]

    m, counts, matched = _spatial_match(f1, f2, lambda p, msg: None)
    assert matched == 2
    assert counts == [2]

    path = build_rows_and_export(f1, f2, m, counts)
    header, body = _read_sheet(path)
    assert len(body) == 2
    ci = header.index(RELATION_COL)
    assert all(r[ci] == RELATION_IN for r in body)


def test_partial_match_unmatched_row_has_empty_l2():
    f1 = [JoinFeature("a", {"小区": "c1"}, Point(116.30, 39.90)),
          JoinFeature("b", {"小区": "c2"}, Point(120.0, 30.0))]
    f2 = [JoinFeature("p1", {"县名": "东城区"}, _square(116.30, 39.90))]

    m, counts, matched = _spatial_match(f1, f2, lambda p, msg: None)
    assert matched == 1

    path = build_rows_and_export(f1, f2, m, counts)
    header, body = _read_sheet(path)
    rels = [r[header.index(RELATION_COL)] for r in body]
    assert RELATION_IN in rels and RELATION_OUT in rels
    # 未匹配行：图层2列空
    oi = rels.index(RELATION_OUT)
    assert body[oi][header.index("县名")] is None


def test_line_anchor_is_bbox_center():
    """线要素按包围盒中心（锚点）判定"""
    line = LineString([(116.0, 39.0), (116.2, 39.2)])
    f1 = [JoinFeature("l", {"编号": "r1"}, line)]
    f2 = [JoinFeature("p", {"县名": "某区"}, _square(116.1, 39.1))]

    m, counts, matched = _spatial_match(f1, f2, lambda p, msg: None)
    assert matched == 1


def test_dangling_polygon_appended():
    f1 = [JoinFeature("a", {"小区": "c1"}, Point(116.30, 39.90))]
    f2 = [JoinFeature("p1", {"县名": "东城区"}, _square(116.30, 39.90)),
          JoinFeature("p2", {"县名": "西城区"}, _square(120.0, 30.0))]

    m, counts, matched = _spatial_match(f1, f2, lambda p, msg: None)
    assert counts == [1, 0]

    path = build_rows_and_export(f1, f2, m, counts)
    header, body = _read_sheet(path)
    rels = [r[header.index(RELATION_COL)] for r in body]
    assert len(body) == 2  # f1 一行 + 零归属面一行
    di = rels.index(RELATION_DANGLING)
    assert body[di][header.index("县名")] == "西城区"
    assert body[di][header.index("小区")] is None


def test_column_name_conflict_gets_prefix():
    f1 = [JoinFeature("a", {"name": "c1"}, Point(116.30, 39.90))]
    f2 = [JoinFeature("p1", {"name": "东城区", "tac": "T1"}, _square(116.30, 39.90))]

    m, counts, matched = _spatial_match(f1, f2, lambda p, msg: None)
    path = build_rows_and_export(f1, f2, m, counts)
    header, _ = _read_sheet(path)
    assert "name" in header
    assert "L2_name" in header
    assert "tac" in header
    assert len(header) == len(set(header))  # 无重复列


# ============================ 主流程 ============================

def test_run_success_stats(monkeypatch):
    f1 = [JoinFeature("a", {"小区": "c1"}, Point(116.30, 39.90)),
          JoinFeature("b", {"小区": "c2"}, Point(120.0, 30.0))]
    f2 = [JoinFeature("p1", {"县名": "东城区"}, _square(116.30, 39.90))]
    _patch_loader(monkeypatch, f1, f2)

    stats = run_attribute_join({"layer1": {}, "layer2": {}}, None, lambda p, msg: None)
    assert stats["total1"] == 2
    assert stats["matched"] == 1
    assert stats["unmatched"] == 1
    assert stats["dangling2"] == 0
    assert Path(stats["exportPath"]).exists()


def test_run_bbox_disjoint_raises(monkeypatch):
    f1 = [JoinFeature("a", {"小区": "c1"}, Point(116.3, 39.9))]
    f2 = [JoinFeature("p1", {"县名": "某县"}, _square(120.0, 30.0))]
    _patch_loader(monkeypatch, f1, f2)

    with pytest.raises(ValueError, match="不存在包含关系"):
        run_attribute_join({"layer1": {}, "layer2": {}}, None, lambda p, msg: None)


def test_run_bbox_overlap_but_point_outside_raises(monkeypatch):
    # L 形面：包围盒含点，但点落在缺口内（面外）
    poly = Polygon([
        (116.29, 39.89), (116.31, 39.89), (116.31, 39.895),
        (116.295, 39.895), (116.295, 39.91), (116.29, 39.91),
    ])
    f1 = [JoinFeature("a", {"小区": "c1"}, Point(116.30, 39.90))]
    f2 = [JoinFeature("p1", {"县名": "某县"}, poly)]
    _patch_loader(monkeypatch, f1, f2)

    with pytest.raises(ValueError, match="不存在包含关系"):
        run_attribute_join({"layer1": {}, "layer2": {}}, None, lambda p, msg: None)


def test_run_empty_layer1_raises(monkeypatch):
    _patch_loader(monkeypatch, [], [JoinFeature("p", {"a": 1}, _square(0, 0))])
    with pytest.raises(ValueError, match="图层1无有效要素"):
        run_attribute_join({"layer1": {}, "layer2": {}}, None, lambda p, msg: None)


def test_run_layer2_not_polygon_raises(monkeypatch):
    _patch_loader(monkeypatch,
                  [JoinFeature("a", {"x": 1}, Point(0, 0))],
                  [JoinFeature("b", {"y": 2}, Point(1, 1))])  # 点要素，过滤后 f2 为空
    with pytest.raises(ValueError, match="面状要素"):
        run_attribute_join({"layer1": {}, "layer2": {}}, None, lambda p, msg: None)


# ===== 经纬度缺失校验与跳过 =====

def test_valid_lonlat_accepts_rejects():
    assert valid_lonlat(116.3, 39.9) is True
    assert valid_lonlat(116.3, -5.0) is True
    # (0,0) = 缺失被填的默认值
    assert valid_lonlat(0, 0) is False
    # None / 空串 / NaN
    assert valid_lonlat(None, 39.9) is False
    assert valid_lonlat("", "") is False
    assert valid_lonlat(float("nan"), 39.9) is False
    # 越界
    assert valid_lonlat(200.0, 30.0) is False
    assert valid_lonlat(116.0, 100.0) is False


def test_run_all_missing_raises(monkeypatch):
    skip = [JoinFeature("b", {"扇区名称": "bad", "扇区ID": "2"},
                        None, REASON_MISSING_COORD)]
    _patch_loader(monkeypatch, [],
                  [JoinFeature("p", {"县名": "x"}, _square(116.3, 39.9))], s1=skip)
    with pytest.raises(ValueError, match="均缺失有效经纬度"):
        run_attribute_join({"layer1": {}, "layer2": {}}, None, lambda p, msg: None)


def test_run_skips_missing_and_reports(monkeypatch):
    valid = [JoinFeature("a", {"扇区名称": "good", "扇区ID": "1"}, Point(116.30, 39.90))]
    skip = [JoinFeature("b", {"扇区名称": "bad", "扇区ID": "2"},
                        None, REASON_MISSING_COORD)]
    _patch_loader(monkeypatch, valid,
                  [JoinFeature("p", {"县名": "某县"}, _square(116.30, 39.90))], s1=skip)

    stats = run_attribute_join({"layer1": {}, "layer2": {}}, None, lambda p, msg: None)
    assert stats["total1"] == 1          # 仅有效要素参与匹配
    assert stats["matched"] == 1
    assert stats["skipped1"] == 1
    assert stats["skipped1Details"] == [
        {"id": "2", "name": "bad", "reason": REASON_MISSING_COORD}
    ]

    header, body = _read_sheet(stats["exportPath"])
    rels = [r[header.index(RELATION_COL)] for r in body]
    assert RELATION_SKIP in rels
    skip_row = body[rels.index(RELATION_SKIP)]
    assert skip_row[header.index("扇区名称")] == "bad"
    assert skip_row[header.index("县名")] is None   # 跳过行图层2列空
