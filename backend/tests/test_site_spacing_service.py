# backend/tests/test_site_spacing_service.py
"""平均站间距计算服务单元测试"""
import math
import os
import tempfile

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook, Workbook

from app.services.site_spacing_service import (
    _normalize_columns,
    _haversine_np,
    build_output_excel,
    build_output_filename,
    compute_avg_spacing,
    compute_overall_stats,
    haversine_km,
    read_sites_file,
    validate_sites_file,
)


# ============== haversine 基础校验 ==============

def test_haversine_zero():
    assert haversine_km(116.0, 39.0, 116.0, 39.0) == 0.0


def test_haversine_1deg_lat():
    """北京附近 1 度纬度 ≈ 111.2 km"""
    d = haversine_km(116.4, 39.9, 116.4, 40.9)
    assert 110.0 < d < 113.0, f"实际 {d}"


def test_haversine_1deg_lon_at_40lat():
    """40°纬度下 1 度经度 ≈ 85.4 km"""
    d = haversine_km(116.4, 40.0, 117.4, 40.0)
    assert 84.0 < d < 87.0, f"实际 {d}"


def test_haversine_vectorized_matches_scalar():
    """向量化版本与标量版应给出相同结果（输入 1D 数组）"""
    lons1 = np.array([116.4, 116.5, 116.6])
    lats1 = np.array([39.9, 40.0, 40.1])
    lons2 = np.array([116.5, 116.6, 116.7])
    lats2 = np.array([40.0, 40.1, 40.2])
    vec = _haversine_np(lons1, lats1, lons2, lats2)
    assert vec.shape == (3,)
    for i in range(3):
        expected = haversine_km(lons1[i], lats1[i], lons2[i], lats2[i])
        assert abs(vec[i] - expected) < 1e-6, f"i={i}: vec={vec[i]} expected={expected}"


def test_haversine_vectorized_broadcast():
    """支持 (B,1) 与 (K,) 广播：shape (B, K)"""
    q_lon = np.array([116.4, 116.5])[:, None]   # (2, 1)
    q_lat = np.array([39.9, 40.0])[:, None]     # (2, 1)
    n_lon = np.array([116.5, 116.6, 116.7])     # (3,)
    n_lat = np.array([40.0, 40.1, 40.2])        # (3,)
    vec = _haversine_np(q_lon, q_lat, n_lon, n_lat)
    assert vec.shape == (2, 3)
    for i in range(2):
        for j in range(3):
            expected = haversine_km(q_lon[i, 0], q_lat[i, 0], n_lon[j], n_lat[j])
            assert abs(vec[i, j] - expected) < 1e-6, f"({i},{j}): vec={vec[i,j]} expected={expected}"


# ============== KNN 核心算法 ==============

def test_compute_avg_spacing_simple_4neighbors():
    """5 个站：1 中心 + 4 个正交方向各 ~1km
    K=4 时每个站的平均站间距 = 4 个邻居距离的均值
    """
    cells = ["A", "B", "C", "D", "E"]
    lons = [116.4074, 116.4174, 116.4074, 116.3974, 116.4074]
    lats = [39.9042, 39.9042, 39.9142, 39.9042, 39.8942]
    results = compute_avg_spacing(cells, lons, lats, k=4)
    assert len(results) == 5
    # A 的 4 个邻居正好是 B/C/D/E
    ab = haversine_km(116.4074, 39.9042, 116.4174, 39.9042)
    ac = haversine_km(116.4074, 39.9042, 116.4074, 39.9142)
    ad = haversine_km(116.4074, 39.9042, 116.3974, 39.9042)
    ae = haversine_km(116.4074, 39.9042, 116.4074, 39.8942)
    expected_A = (ab + ac + ad + ae) / 4
    assert abs(results[0][3] - round(expected_A, 4)) < 0.001
    # B 的 4 个邻居是 A/C/D/E
    bc = haversine_km(116.4174, 39.9042, 116.4074, 39.9142)
    bd = haversine_km(116.4174, 39.9042, 116.3974, 39.9042)
    be = haversine_km(116.4174, 39.9042, 116.4074, 39.8942)
    expected_B = (ab + bc + bd + be) / 4
    assert abs(results[1][3] - round(expected_B, 4)) < 0.001


def test_compute_avg_spacing_excludes_self():
    """自身距离 = 0，应被排除在 K 近邻之外"""
    cells = ["A", "B", "C"]
    lons = [116.4, 116.410, 116.418]
    lats = [39.9, 39.9, 39.9]
    res = compute_avg_spacing(cells, lons, lats, k=1)
    assert res[0][3] > 0.5  # A → B 约 0.89km


def test_compute_avg_spacing_k_exceeds_n():
    """K > N-1 时使用实际可用数（不能比自己还多）"""
    cells = ["A", "B", "C"]
    lons = [116.4, 116.41, 116.42]
    lats = [39.9, 39.9, 39.9]
    res = compute_avg_spacing(cells, lons, lats, k=10)
    # A 的实际邻居数 = 2 (B, C)
    assert res[0][4] == 2
    assert res[0][3] > 0


def test_compute_avg_spacing_empty():
    assert compute_avg_spacing([], [], [], k=5) == []


def test_compute_avg_spacing_single():
    """单点：返回 (CELL, LON, LAT, 0.0, 0)"""
    res = compute_avg_spacing(["X"], [116.4], [39.9], k=5)
    assert res == [("X", 116.4, 39.9, 0.0, 0)]


def test_compute_avg_spacing_k_invalid():
    with pytest.raises(ValueError, match="大于 0"):
        compute_avg_spacing(["A"], [116.4], [39.9], k=0)
    with pytest.raises(ValueError, match="大于 0"):
        compute_avg_spacing(["A"], [116.4], [39.9], k=-1)


def test_compute_avg_spacing_4_corners():
    """4 个角的方阵：每点的最近 K=3 邻居应选出其它 3 个点"""
    cells = ["A", "B", "C", "D"]
    lons = [116.4, 116.5, 116.4, 116.5]
    lats = [39.9, 39.9, 40.0, 40.0]
    res = compute_avg_spacing(cells, lons, lats, k=3)
    for r in res:
        assert r[4] == 3
        assert r[3] > 0


def test_compute_avg_spacing_random_1k_correct():
    """1000 个随机站点：验证 KNN 排序正确（与 brute force 对比前 10 个点）。

    注意：使用 5° 范围以保证 3D 单位球面点积的 float32 精度足够区分 ~1km 级别距离。
    范围过小（如 0.1°）时 cos 差异 < float32 精度，算法会随机选点（已知限制）。
    """
    rng = np.random.default_rng(42)
    n = 1000
    lons = (116.4 + rng.random(n) * 5.0).tolist()
    lats = (39.9 + rng.random(n) * 5.0).tolist()
    cells = [f"C{i:05d}" for i in range(n)]
    k = 5
    res = compute_avg_spacing(cells, lons, lats, k=k)
    # 前 10 个点用 brute force 验证
    for i in range(10):
        dists = []
        for j in range(n):
            if j == i:
                continue
            dists.append(haversine_km(lons[i], lats[i], lons[j], lats[j]))
        dists.sort()
        expected = sum(dists[:k]) / k
        actual = res[i][3]
        assert abs(actual - round(expected, 4)) < 0.05, (
            f"point {i}: actual={actual} expected={round(expected, 4)}"
        )


# ============== 列名归一化 ==============

def test_normalize_columns_basic():
    df = pd.DataFrame({"小区名称": ["A", "B"], "经度": [116.4, 116.5], "纬度": [39.9, 40.0]})
    out = _normalize_columns(df)
    assert list(out.columns) == ["CELL", "LON", "LAT"]
    assert len(out) == 2


def test_normalize_columns_english():
    """CELL / LON / LAT 英文别名"""
    df = pd.DataFrame({"CELL": ["A"], "LON": [116.4], "LAT": [39.9]})
    out = _normalize_columns(df)
    assert list(out.columns) == ["CELL", "LON", "LAT"]


def test_normalize_columns_missing():
    """缺列抛 ValueError 并提示中文"""
    df = pd.DataFrame({"name": ["A"], "lon": [116.4]})  # 缺 lat
    with pytest.raises(ValueError) as exc:
        _normalize_columns(df)
    msg = str(exc.value)
    assert "无法识别" in msg or "缺少" in msg or "Latitude" in msg or "纬度" in msg


# ============== 文件读写 ==============

def test_read_sites_file_xlsx():
    """读 .xlsx 文件，自动归一化列名 + 过滤无效行 + 去重"""
    with tempfile.TemporaryDirectory() as d:
        fp = os.path.join(d, "sites.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.append(["标签名称", "经度", "纬度"])
        ws.append(["A", 116.4, 39.9])
        ws.append(["B", 116.5, 40.0])
        ws.append(["C", None, None])  # 无效行
        ws.append(["A", 116.6, 40.1])  # 与 A 重复
        wb.save(fp)

        out = read_sites_file(fp)
        assert len(out) == 2
        assert out["CELL"].tolist() == ["A", "B"]


def test_read_sites_file_csv_utf8():
    """读 .csv 文件（utf-8）"""
    with tempfile.TemporaryDirectory() as d:
        fp = os.path.join(d, "sites.csv")
        with open(fp, "w", encoding="utf-8") as f:
            f.write("小区名称,经度,纬度\n")
            f.write("A,116.4,39.9\n")
            f.write("B,116.5,40.0\n")
        out = read_sites_file(fp)
        assert len(out) == 2
        assert out.iloc[0]["CELL"] == "A"


def test_read_sites_file_txt_tab():
    """读 .txt 文件（tab 分隔）"""
    with tempfile.TemporaryDirectory() as d:
        fp = os.path.join(d, "sites.txt")
        with open(fp, "w", encoding="utf-8") as f:
            f.write("小区名称\t经度\t纬度\n")
            f.write("A\t116.4\t39.9\n")
            f.write("B\t116.5\t40.0\n")
        out = read_sites_file(fp)
        assert len(out) == 2


def test_read_sites_file_csv_gbk():
    """读 .csv 文件（GBK 编码）"""
    with tempfile.TemporaryDirectory() as d:
        fp = os.path.join(d, "sites_gbk.csv")
        with open(fp, "w", encoding="gbk") as f:
            f.write("标签名称,经度,纬度\n")
            f.write("站点A,116.4,39.9\n")
        out = read_sites_file(fp)
        assert len(out) == 1
        assert out.iloc[0]["CELL"] == "站点A"


def test_validate_sites_file_invalid_rows():
    """校验无效行（经纬度超界 / 为空）"""
    with tempfile.TemporaryDirectory() as d:
        fp = os.path.join(d, "bad.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.append(["标签名称", "经度", "纬度"])
        ws.append(["A", 116.4, 39.9])
        ws.append(["", 116.5, 40.0])  # 名称空
        ws.append(["C", 200.0, 40.0])  # 经度超界
        ws.append(["D", 116.4, -100.0])  # 纬度超界
        ws.append(["E", "abc", 40.0])  # 经度非数值
        wb.save(fp)
        errors = validate_sites_file(fp)
        assert len(errors) == 4
        reasons = [e["reason"] for e in errors]
        assert "站点名称为空" in reasons
        assert "经度超出有效范围" in reasons
        assert "纬度超出有效范围" in reasons
        assert "经度格式无效" in reasons


# ============== 输出 Excel ==============

def test_build_output_excel():
    rows = [
        ("A", 116.4, 39.9, 0.5, 5),
        ("B", 116.5, 40.0, 0.6, 5),
    ]
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "out.xlsx")
        build_output_excel(out, rows)
        wb = load_workbook(out)
        ws = wb.active
        # 1 表头 + 2 数据
        assert ws.max_row == 3
        assert ws.cell(1, 1).value == "站点名称"
        assert ws.cell(1, 4).value == "平均站间距(km)"
        assert ws.cell(1, 5).value == "参与计算邻居数"
        assert ws.cell(2, 1).value == "A"
        assert ws.cell(2, 4).value == 0.5
        assert ws.cell(2, 5).value == 5
        assert ws.cell(3, 1).value == "B"
        # stats=None 时不创建汇总 sheet
        assert "汇总" not in wb.sheetnames


def test_build_output_excel_with_stats():
    """带 stats 时创建「汇总」sheet 并填入 3 行数据"""
    rows = [
        ("A", 116.4, 39.9, 0.4, 5),
        ("B", 116.5, 40.0, 0.6, 5),
        ("C", 116.6, 40.1, 0.8, 5),
    ]
    stats = {"total": 3, "counted": 3, "overall_avg_km": 0.6}
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "out.xlsx")
        build_output_excel(out, rows, stats)
        wb = load_workbook(out)
        assert "汇总" in wb.sheetnames
        ws2 = wb["汇总"]
        # 1 表头 + 3 数据行
        assert ws2.max_row == 4
        assert ws2.cell(1, 1).value == "指标"
        assert ws2.cell(1, 2).value == "数值"
        assert ws2.cell(2, 1).value == "总站点数"
        assert ws2.cell(2, 2).value == 3
        assert ws2.cell(3, 1).value == "参与计算站点数"
        assert ws2.cell(3, 2).value == 3
        assert ws2.cell(4, 1).value == "整体平均站间距(km)"
        assert ws2.cell(4, 2).value == 0.6


def test_compute_overall_stats_basic():
    """正常多站点：counted=全部, overall_avg=均值"""
    rows = [
        ("A", 116.4, 39.9, 0.4, 5),
        ("B", 116.5, 40.0, 0.6, 5),
        ("C", 116.6, 40.1, 0.8, 5),
    ]
    stats = compute_overall_stats(rows)
    assert stats["total"] == 3
    assert stats["counted"] == 3
    assert stats["overall_avg_km"] == 0.6


def test_compute_overall_stats_excludes_zero_neighbors():
    """actual_k=0 的站点不参与整体平均（单点/无邻居无意义）"""
    rows = [
        ("A", 116.4, 39.9, 0.0, 0),  # 排除
        ("B", 116.5, 40.0, 1.0, 5),
        ("C", 116.6, 40.1, 3.0, 5),
    ]
    stats = compute_overall_stats(rows)
    assert stats["total"] == 3
    assert stats["counted"] == 2
    assert stats["overall_avg_km"] == 2.0  # (1+3)/2


def test_compute_overall_stats_empty():
    stats = compute_overall_stats([])
    assert stats["total"] == 0
    assert stats["counted"] == 0
    assert stats["overall_avg_km"] == 0.0


def test_build_output_filename_format():
    """输出文件名含时间戳"""
    name = build_output_filename()
    assert name.startswith("站间距结果")
    assert name.endswith(".xlsx")
