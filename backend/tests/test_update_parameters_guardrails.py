"""工参更新（现网工参合并进全量工参）：失败必须显式，不允许"提示完成但没写入"。

背景（2026-10-08 用户报）：桌面端执行工参更新提示完成，但下载的文件与更新前一致。
根因是现网工参 CSV 在本机缺失/选错数据集时，合并会静默跑空并照样产出"新文件"。
这里锁定三条契约：
1. 现网工参数据体缺失 → 明确报错（并提示先同步/重新上传），不产出假成功；
2. 一个小区都没匹配上、也没新增行 → 报错（提示核对是否为同一批数据）；
3. 正常合并 → 新文件确实写入现网值，且结果里带统计（匹配/新增行数）。
"""
from __future__ import annotations

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.models.database import Base
from app.services.data_service import data_service

FULL_ID = "test-upd-full"
CUR_ID = "test-upd-current"
LTE_HEADER = ["eNodeB标识", "小区标识", "小区名称", "基站名称", "经度", "纬度"]


@pytest.fixture()
def update_env(tmp_path, monkeypatch):
    """临时库 + 临时全量工参 Excel + 现网工参目录（CSV 可选）。"""
    from app.core.config import settings

    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(bind=engine)
    Session = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    monkeypatch.setattr("app.services.data_service.SessionLocal", Session)

    data_service.index = {}

    full_dir = settings.DATA_DIR / FULL_ID
    full_dir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LTE Project Parameters"
    ws.append(LTE_HEADER)
    ws.append([])  # 标题/副标题行（解析契约：数据从第 4 行开始）
    ws.append([])
    rows = [("1001", "1", "A-1"), ("1001", "2", "A-2"), ("1002", "1", "B-1")]
    for enb, cell, name in rows:
        ws.append([enb, cell, name, None, 113.0, 23.0])
    wb.save(full_dir / "original.xlsx")
    wb.close()

    data_service.index[FULL_ID] = {
        "id": FULL_ID, "name": "ProjectParameter_test.xlsx", "type": "excel",
        "fileType": "full_params", "status": "ready", "user_id": 1,
    }
    cur_dir = settings.DATA_DIR / CUR_ID
    cur_dir.mkdir(parents=True, exist_ok=True)
    data_service.index[CUR_ID] = {
        "id": CUR_ID, "name": "BaselineLab_test.zip", "type": "map",
        "subType": "offline", "status": "ready", "user_id": 1,
    }
    return {"current_dir": cur_dir, "settings": settings}


def _write_current_csv(current_dir, rows):
    import pandas as pd

    df = pd.DataFrame(rows, columns=["eNBId", "cellLocalId", "eNBName"])
    path = current_dir / "LTE_SDR_CellInfo_20260101000000.csv"
    df.to_csv(path, index=False)
    return path


def _new_excel_values(new_file_id, *, col: str):
    from app.core.config import settings

    path = settings.DATA_DIR / new_file_id / "original.xlsx"
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        ws = wb["LTE Project Parameters"]
        header = [str(c).strip().split("\n")[0] if c else "" for c in
                  next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        idx = header.index(col)
        return [row[idx] for row in ws.iter_rows(min_row=4, values_only=True)]
    finally:
        wb.close()


def test_missing_current_params_csv_fails_loudly(update_env):
    """现网工参目录里没有 CSV（本机镜像不完整/选错数据集）→ 必须报错，不产出新文件。"""
    with pytest.raises(ValueError, match="现网工参数据不完整"):
        data_service.update_parameters(FULL_ID, CUR_ID, username="u", user_id=1)


def test_unparseable_current_csv_fails_loudly(update_env):
    """CSV 存在但不是网管 CellInfo 结构（无法解析出任何记录）→ 报错而不是产出未变化的文件。"""
    import pandas as pd

    # 缺少 eNBId/cellLocalId 列 → _load_current_params 会整份跳过
    pd.DataFrame([(1, "A")], columns=["cgi", "name"]).to_csv(
        update_env["current_dir"] / "LTE_SDR_CellInfo_20260101000000.csv", index=False
    )
    with pytest.raises(ValueError, match="未匹配到任何小区"):
        data_service.update_parameters(FULL_ID, CUR_ID, username="u", user_id=1)


def test_unmatched_cells_are_appended(update_env):
    """现网有全量工参里没有的小区 → 追加新行（既有语义，不能被守卫误伤）。"""
    _write_current_csv(update_env["current_dir"], [
        ("1001", "1", "现网-站1"),
        ("1003", "1", "现网-新增站"),
    ])

    result = data_service.update_parameters(FULL_ID, CUR_ID, username="u", user_id=1)

    stats = result["stats"]["LTE"]
    assert stats["matched"] == 1 and stats["appended"] == 1
    names = _new_excel_values(result["newFileId"], col="基站名称")
    assert "现网-新增站" in names


def test_merge_writes_values_and_reports_stats(update_env):
    """正常合并：现网值写入新 Excel，返回统计（匹配/新增行数）。"""
    _write_current_csv(update_env["current_dir"], [
        ("1001", "1", "现网-站1"),
        ("1001", "2", "现网-站2"),
        ("1002", "1", "现网-站3"),
    ])

    result = data_service.update_parameters(FULL_ID, CUR_ID, username="u", user_id=1)

    assert result["newFileId"] and result["newFileName"].endswith(".xlsx")
    stats = result["stats"]["LTE"]
    assert stats["matched"] == 3 and stats["appended"] == 0

    names = _new_excel_values(result["newFileId"], col="基站名称")
    assert names[:3] == ["现网-站1", "现网-站2", "现网-站3"]
    # 新条目已注册、旧全量工参已替换
    assert result["newFileId"] in data_service.index
    assert FULL_ID not in data_service.index


# ---------------------------------------------------------------- 桌面端补拉

def test_repair_dataset_repulls_missing_mirror_in_shared_mode(tmp_path, monkeypatch):
    """桌面共享模式：本机镜像不完整时按服务器修订号补拉一次（工参更新读不到 CSV 的自愈路径）。"""
    from types import SimpleNamespace

    from app.core.config import settings
    from app.services import desktop_cloud_client as cloud
    from app.services import desktop_sync_service as sync

    monkeypatch.setattr(settings, "NPT_DESKTOP", True)
    monkeypatch.setattr(settings, "DATA_ROOT", tmp_path)
    monkeypatch.setattr(sync, "state_file", lambda: tmp_path / "sync_state.json")
    sync.save_state(run_mode="shared")

    seen = {}
    monkeypatch.setattr(sync, "_require_active_license", lambda: None)
    monkeypatch.setattr(
        sync, "_local_group",
        lambda: (SimpleNamespace(id=1, username="u"), "河源", "电信"),
    )
    monkeypatch.setattr(cloud, "sync_manifest", lambda: {"items": [{"data_id": "d1", "rev": 7}]})
    monkeypatch.setattr(
        sync, "_pull",
        lambda data_id, user, city, operator, rev, server_updated_at=None: seen.update(
            data_id=data_id, rev=rev, server_updated_at=server_updated_at
        )
        or {"data_id": data_id, "action": "pull", "rev": rev},
    )

    out = sync.repair_dataset("d1")
    assert out["repaired"] is True
    # 补拉同时把服务器最后更新时间转发给 _pull（LWW 基线）
    assert seen == {"data_id": "d1", "rev": 7, "server_updated_at": None}


def test_repair_dataset_skips_when_local_dirty(tmp_path, monkeypatch):
    """本机有未上行修改时不得补拉（避免覆盖用户数据）。"""
    from types import SimpleNamespace

    from app.core.config import settings
    from app.services import desktop_sync_service as sync

    monkeypatch.setattr(settings, "NPT_DESKTOP", True)
    monkeypatch.setattr(settings, "DATA_ROOT", tmp_path)
    monkeypatch.setattr(sync, "state_file", lambda: tmp_path / "sync_state.json")
    sync.save_state(run_mode="shared", datasets={"d1": {"rev": 1, "dirty": True}})
    monkeypatch.setattr(sync, "_local_group",
                        lambda: (SimpleNamespace(id=1, username="u"), "河源", "电信"))
    monkeypatch.setattr(sync, "_require_active_license", lambda: None)
    monkeypatch.setattr(
        sync, "_pull",
        lambda *a, **k: pytest.fail("dirty 数据集不得被补拉覆盖"),
    )

    assert sync.repair_dataset("d1")["repaired"] is False


def test_repair_dataset_noop_outside_desktop_shared():
    """网页端/单机模式不做补拉（避免在服务器上误删重装数据集）。"""
    from app.core.config import settings
    from app.services import desktop_sync_service as sync

    assert settings.NPT_DESKTOP is False
    assert sync.repair_dataset("d1")["repaired"] is False
