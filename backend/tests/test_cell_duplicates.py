"""制式分离查重与新站录入重复消息测试"""
import pytest

from app.services.data_service import data_service


def test_duplicate_message_same_type():
    """同制式重复 → '已重复，信息已更新覆盖'"""
    assert data_service._duplicate_message("NR", "NR", "X") == "[NR]小区 X 已重复，信息已更新覆盖"
    assert data_service._duplicate_message("LTE", "LTE", "Y") == "[LTE]小区 Y 已重复，信息已更新覆盖"
    assert data_service._duplicate_message("LTE", None, "Z") == "[LTE]小区 Z 已重复，信息已更新覆盖"
    assert data_service._duplicate_message("LTE", "", "E") == "[LTE]小区 E 已重复，信息已更新覆盖"


def test_duplicate_message_cross_type():
    """跨制式重复 → '已在[其他制式]工参中存在，请检查'"""
    assert data_service._duplicate_message("LTE", "NR", "X") == "[LTE]小区 X 已在[NR]工参中存在，请检查"
    assert data_service._duplicate_message("NR", "LTE", "Y") == "[NR]小区 Y 已在[LTE]工参中存在，请检查"


@pytest.fixture()
def dup_env():
    """构造含 LTE/NR 小区的全量工参环境（纯 DB fixture，供后续任务复用）"""
    import uuid
    from app.core.database import SessionLocal
    from app.models.database import Workbook, Site, Sector

    db = SessionLocal()
    wb_id = f"test-{uuid.uuid4()}"
    db.add(Workbook(
        id=wb_id, user_id=1, name="dup-test", file_type="full_params",
        status="completed", created_at="2026-07-31",
    ))
    # 先 commit Workbook，再 flush Site：
    # models 未声明 relationship()，SQLAlchemy 无法按 FK 依赖排序 insert，
    # 若在同一个 flush 中同时插入 Workbook 与 Site，Site 会先被写入而违反外键。
    db.commit()
    site_lte = Site(
        workbook_id=wb_id, user_id=1, site_id="111", name="LTE站",
        network_type="LTE", longitude=1.0, latitude=1.0,
    )
    db.add(site_lte)
    db.flush()
    db.add(Sector(
        site_id_fk=site_lte.id, workbook_id=wb_id, user_id=1,
        sector_id="111_1", name="LTE小区", longitude=1.0, latitude=1.0, azimuth=0,
    ))
    site_nr = Site(
        workbook_id=wb_id, user_id=1, site_id="222", name="NR站",
        network_type="NR", longitude=2.0, latitude=2.0,
        mcc="460", mnc="11",
    )
    db.add(site_nr)
    db.flush()
    # NR 电信小区（460-11）；共享小区的联通 PLMN（460-01）是另一条 Sector，见下方用例
    db.add(Sector(
        site_id_fk=site_nr.id, workbook_id=wb_id, user_id=1,
        sector_id="222_1", name="NR小区", longitude=2.0, latitude=2.0, azimuth=0,
        mcc="460", mnc="11",
    ))
    db.commit()

    # 注入 index，让 data_service 能定位 full_params 文件。
    # 实现按插入顺序取第一个 full_params，若 index 中已存在真实 full_params 条目
    # 会先于本 fixture 的 wb_id 被选中，导致查重针对错误的 workbook。
    # 因此临时移除既有 full_params 条目，仅保留本 fixture 注入项，teardown 时恢复。
    saved_index = dict(data_service.index)
    saved_full_params = [
        (k, v) for k, v in saved_index.items()
        if isinstance(v, dict) and v.get("fileType") == "full_params"
    ]
    for k, _ in saved_full_params:
        data_service.index.pop(k)
    data_service.index[wb_id] = {"fileType": "full_params"}
    db.close()

    yield wb_id

    db = SessionLocal()
    try:
        for s in db.query(Sector).filter(Sector.workbook_id == wb_id).all():
            db.delete(s)
        for s in db.query(Site).filter(Site.workbook_id == wb_id).all():
            db.delete(s)
        wb = db.query(Workbook).filter(Workbook.id == wb_id).first()
        if wb:
            db.delete(wb)
        db.commit()
    finally:
        db.close()
        # 无论 DB 清理是否成功，都必须恢复共享单例 index，
        # 否则 fixture 的 wb_id 会永久残留在 index 中污染后续测试。
        data_service.index.clear()
        data_service.index.update(saved_index)


def test_get_sector_network_type(dup_env):
    """按 (site_id, cell_id) 查询已有小区制式"""
    assert data_service._get_sector_network_type(dup_env, "111", "1") == "LTE"
    assert data_service._get_sector_network_type(dup_env, "222", "1") == "NR"
    assert data_service._get_sector_network_type(dup_env, "999", "9") is None
    # 空 siteId/cellId → 早返回 None，不触发 DB 查询
    assert data_service._get_sector_network_type(dup_env, "", "1") is None
    assert data_service._get_sector_network_type(dup_env, "111", "") is None


def test_get_sector_network_type_pure_cell_id(dup_env):
    """纯 cellId 存储格式（如 sector_id='9'）也能被制式查询匹配到"""
    from app.core.database import SessionLocal
    from app.models.database import Site, Sector

    db = SessionLocal()
    site = None
    sector = None
    try:
        site = Site(
            workbook_id=dup_env, user_id=1, site_id="333", name="纯格式站",
            network_type="NR", longitude=3.0, latitude=3.0,
        )
        db.add(site)
        db.flush()
        # 纯小区标识，不带 "siteId_" 前缀（全量工参解析写入的存储格式）
        sector = Sector(
            site_id_fk=site.id, workbook_id=dup_env, user_id=1,
            sector_id="9", name="纯格式小区", longitude=3.0, latitude=3.0, azimuth=0,
        )
        db.add(sector)
        db.commit()

        # 纯格式 cellId 也能通过 sector_id.in_(["9", "333_9"]) 匹配到
        assert data_service._get_sector_network_type(dup_env, "333", "9") == "NR"
    finally:
        if sector is not None:
            db.delete(sector)
        if site is not None:
            db.delete(site)
        db.commit()
        db.close()


def test_batch_duplicate_message_same_and_cross(dup_env, tmp_path, monkeypatch):
    """批量查重:同制式→已重复覆盖,跨制式→已在[其他制式]工参中存在,请检查"""
    # 纯消息层验证（不依赖文件系统）
    assert data_service._duplicate_message("LTE", "LTE", "LTE小区") == "[LTE]小区 LTE小区 已重复，信息已更新覆盖"
    assert data_service._duplicate_message("LTE", "NR", "NR小区") == "[LTE]小区 NR小区 已在[NR]工参中存在，请检查"

    # 集成层：构造内存 Excel + 隔离 DATA_DIR，走 batch_add_cells_to_full_params 全流程
    from openpyxl import Workbook
    from app.core.config import settings
    from app.core.database import SessionLocal
    from app.models.database import Site, Sector

    wb_id = dup_env  # 库中已有 LTE 111_1 + NR 222_1

    # 隔离文件系统：Excel 原始文件与 index 落到 tmp_path，不污染真实 data/ 目录
    monkeypatch.setattr(settings, "DATA_DIR", tmp_path)
    monkeypatch.setattr(data_service, "data_index_file", tmp_path / "index.json")

    excel_path = tmp_path / wb_id / "original.xlsx"
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    lte_ws = wb.active
    lte_ws.title = "LTE Project Parameters"
    lte_ws.append(["小区经度", "小区纬度", "小区名称", "系统制式"])
    nr_ws = wb.create_sheet("NR Project Parameters")
    nr_ws.append(["小区经度", "小区纬度", "小区名称", "系统制式"])
    wb.save(excel_path)
    wb.close()

    cells = [
        {"networkType": "LTE", "siteId": "111", "cellId": "1", "cellName": "LTE小区"},  # 同制式(LTE)重复
        {"networkType": "LTE", "siteId": "222", "cellId": "1", "cellName": "NR小区"},  # 跨制式(NR)重复
        {"networkType": "NR", "siteId": "999", "cellId": "9", "cellName": "新小区",
         "mcc": "460", "mnc": "11", "userId": 1},  # 新小区
        # NR 共享小区：库内 222_1 仅有电信 460-11
        {"networkType": "NR", "siteId": "222", "cellId": "1", "cellName": "NR小区-联通",
         "mcc": "460", "mnc": "01", "userId": 1},  # 异 MNC（联通 460-01）→ 放行，新导入
        {"networkType": "NR", "siteId": "222", "cellId": "1", "cellName": "NR小区-电信重复",
         "mcc": "460", "mnc": "11", "userId": 1},  # 同 PLMN（电信 460-11）→ 同制式重复
    ]
    result = data_service.batch_add_cells_to_full_params(cells)

    r0, r1, r2, r3, r4 = result["results"]
    assert r0["message"] == "[LTE]小区 LTE小区 已重复，信息已更新覆盖"
    assert r1["message"] == "[LTE]小区 NR小区 已在[NR]工参中存在，请检查"
    assert r2["message"] == "新增小区 新小区 成功"
    assert r3["message"] == "新增小区 NR小区-联通 成功"          # 共享小区异 MNC 允许录入
    assert r4["message"] == "[NR]小区 NR小区-电信重复 已重复，信息已更新覆盖"  # 同 PLMN 重复拦截
    assert all(r["success"] for r in (r0, r1, r2, r3, r4))

    # 落库校验：重复行不导入；联通共享小区作为第二条 PLMN Sector 落库
    db = SessionLocal()
    try:
        def _count(site_id, sector_id, mnc=None):
            q = db.query(Sector).join(Site, Sector.site_id_fk == Site.id).filter(
                Site.workbook_id == wb_id,
                Site.site_id == site_id,
                Sector.sector_id == sector_id,
            )
            if mnc is not None:
                q = q.filter(Sector.mnc == mnc)
            return q.count()
        assert _count("111", "111_1") == 1          # 同制式重复未重复导入
        assert _count("222", "222_1", "11") == 1    # 原电信 PLMN 保持 1 条
        assert _count("222", "222_1", "01") == 1    # 联通共享 PLMN 新导入 1 条
        assert _count("999", "999_9") == 1          # 新小区成功落库
    finally:
        db.close()


def test_check_duplicates_lte_flags(dup_env):
    """LTE 行按 基站_小区 键：同制式(LTE)→lte，同 基站_小区 存在 NR→nr（跨制式）"""
    result = data_service.check_cell_duplicates([
        {"siteId": "111", "cellId": "1", "networkType": "LTE"},   # 库中为 LTE → lte=True
        {"siteId": "222", "cellId": "1", "networkType": "LTE"},   # 库中为 NR  → nr=True（跨制式）
        {"siteId": "999", "cellId": "9", "networkType": "LTE"},   # 不存在
        {"siteId": "", "cellId": "", "networkType": "LTE"},       # 空 siteId/cellId → 跳过，不产生 key
    ])
    assert result["exists"]["111_1"] == {"lte": True, "nr": False}
    assert result["exists"]["222_1"] == {"lte": False, "nr": True}
    assert result["exists"]["999_9"] == {"lte": False, "nr": False}
    assert "111_1" in result["duplicates"]
    assert "222_1" in result["duplicates"]
    assert "999_9" not in result["duplicates"]
    assert "" not in result["exists"]


def test_check_duplicates_nr_plmn_aware(dup_env):
    """NR 行按 MCC_MNC_基站_小区 键：同 PLMN→nr(同制式)；异 MNC 共享小区放行；同 基站_小区 为 LTE→lte(跨制式)"""
    result = data_service.check_cell_duplicates([
        {"siteId": "222", "cellId": "1", "networkType": "NR", "mcc": "460", "mnc": "11"},  # 同 PLMN(电信) → nr=True
        {"siteId": "222", "cellId": "1", "networkType": "NR", "mcc": "460", "mnc": "01"},  # 联通共享(异 MNC) → 放行
        {"siteId": "111", "cellId": "1", "networkType": "NR", "mcc": "460", "mnc": "11"},  # 同 基站_小区 为 LTE → lte=True
        {"siteId": "999", "cellId": "9", "networkType": "NR", "mcc": "460", "mnc": "11"},  # 不存在
    ])
    # MNC '01' 经 int 归一化为 '1'，与库内 mnc='11' 不是同一 PLMN
    assert result["exists"]["460_11_222_1"] == {"lte": False, "nr": True}
    assert result["exists"]["460_1_222_1"] == {"lte": False, "nr": False}
    assert result["exists"]["460_11_111_1"] == {"lte": True, "nr": False}
    assert result["exists"]["460_11_999_9"] == {"lte": False, "nr": False}
    assert "460_11_222_1" in result["duplicates"]
    assert "460_1_222_1" not in result["duplicates"]
    assert "460_11_111_1" in result["duplicates"]


def test_check_duplicates_endpoint_smoke(dup_env):
    """check-duplicates 端点用新 schema 返回双制式结构（端到端冒烟）"""
    from fastapi.testclient import TestClient
    from app.api import create_app
    from app.core.database import SessionLocal, init_db
    from app.models.database import User
    from app.services.auth_service import auth_service

    # 确保 users 表存在（不依赖 app lifespan）
    init_db()

    username = "dup_smoke_user"
    db = SessionLocal()
    try:
        # 构造具备 can_cell_entry 权限的测试用户 + 合法 token。
        # 该 token 未创建 UserSession 记录 → is_token_revoked 与 is_session_idle_expired 均放行。
        existing = auth_service.get_user_by_username(db, username)
        if existing:
            db.delete(existing)
            db.commit()
        user, err = auth_service.create_user(db, username, "Test@12345")
        assert err is None, f"创建测试用户失败: {err}"
        user.can_cell_entry = True
        db.commit()
        db.refresh(user)
        token, _ = auth_service.create_access_token(user.id, user.username, False)
    finally:
        db.close()

    # 不使用 with 语句创建 TestClient：跳过 lifespan，避免 cleanup_index()
    # 清掉 dup_env fixture 注入的全量工参 index 条目，导致查重找不到目标文件。
    try:
        client = TestClient(create_app())
        resp = client.post(
            "/api/v1/data/full-params/cells/check-duplicates",
            headers={"Authorization": f"Bearer {token}"},
            json={
                "cells": [
                    {"siteId": "111", "cellId": "1", "networkType": "NR", "mcc": "460", "mnc": "11"},  # 库中 LTE → lte=True
                    {"siteId": "222", "cellId": "1", "networkType": "NR", "mcc": "460", "mnc": "11"},  # 同 PLMN NR → nr=True
                    {"siteId": "222", "cellId": "1", "networkType": "NR", "mcc": "460", "mnc": "01"},  # 联通共享 → 放行
                    {"siteId": "999", "cellId": "9", "networkType": "NR", "mcc": "460", "mnc": "11"},  # 不存在
                ]
            },
        )
        assert resp.status_code == 200, f"HTTP {resp.status_code}: {resp.text}"
        body = resp.json()
        assert body["success"] is True
        data = body["data"]
        assert data["exists"]["460_11_111_1"] == {"lte": True, "nr": False}
        assert data["exists"]["460_11_222_1"] == {"lte": False, "nr": True}
        assert data["exists"]["460_1_222_1"] == {"lte": False, "nr": False}
        assert data["exists"]["460_11_999_9"] == {"lte": False, "nr": False}
        assert "460_11_111_1" in data["duplicates"]
        assert "460_11_222_1" in data["duplicates"]
        assert "460_1_222_1" not in data["duplicates"]
        assert "460_11_999_9" not in data["duplicates"]
    finally:
        # 无论断言是否通过，都清理测试用户，避免泄漏用户行
        db = SessionLocal()
        try:
            user = auth_service.get_user_by_username(db, username)
            if user:
                db.delete(user)
                db.commit()
        finally:
            db.close()


def _fp_df(rows):
    """构造全量工参 DataFrame（含 LTE/NR 常用列的超集），rows 为每行覆盖字段 dict"""
    import pandas as pd

    default = {
        "eNodeB标识": "0", "gNodeB标识": "0", "小区标识": "0", "小区名称": "X",
        "基站名称": "站", "移动国家码": "460", "移动网络码": "11",
        "基站经度": 113.0, "基站纬度": 23.0, "天线方向角": 0, "天线挂高": 30,
        "物理小区识别码": 0, "跟踪区码": 1, "小区覆盖类型": 1, "是否共享": "否",
        "第一分组": "电信中兴", "下行链路的中心载频": 1825.0, "填写SSB频点": 327040.0,
        "管理网元ID": "2001",
    }
    return pd.DataFrame([{**default, **r} for r in rows])


def test_full_params_dup_report_lte():
    """全量工参解析：LTE 按 基站_小区 检测重复行并生成报告"""
    data_service.full_params_dup_report = {}
    df = _fp_df([
        {"eNodeB标识": "100", "小区标识": "1", "小区名称": "A", "天线方向角": 0},
        {"eNodeB标识": "100", "小区标识": "1", "小区名称": "A重复", "天线方向角": 30},
        {"eNodeB标识": "200", "小区标识": "2", "小区名称": "B", "基站经度": 114.0,
         "基站纬度": 24.0, "天线方向角": 60},
    ])
    sites = data_service._parse_full_params_dataframe(df, "LTE")
    assert sum(len(s.get("sectors", [])) for s in sites) == 2
    rep = data_service.full_params_dup_report.get("LTE")
    assert rep is not None
    assert rep["groups"] == 1 and rep["extra"] == 1
    assert rep["samples"][0]["siteId"] == "100"
    assert rep["samples"][0]["cellId"] == "1"


def test_full_params_dup_report_nr_plmn_aware():
    """NR 同基站/同小区不同 MNC 是两条 PLMN，不算重复；同 PLMN 重复才计入"""
    data_service.full_params_dup_report = {}
    df = _fp_df([
        {"gNodeB标识": "300", "小区标识": "1", "移动网络码": "11", "小区名称": "T"},
        {"gNodeB标识": "300", "小区标识": "1", "移动网络码": "01", "小区名称": "U"},   # 联通共享，非重复
        {"gNodeB标识": "300", "小区标识": "1", "移动网络码": "11", "小区名称": "T重复"},  # 同 PLMN 重复
    ])
    sites = data_service._parse_full_params_dataframe(df, "NR")
    # 去重后 300 站保留 2 个扇区（电信/联通两条 PLMN）
    assert sum(len(s.get("sectors", [])) for s in sites) == 2
    rep = data_service.full_params_dup_report.get("NR")
    assert rep is not None
    assert rep["groups"] == 1 and rep["extra"] == 1
    assert rep["samples"][0]["mnc"] == "11"


def test_full_params_no_duplicates_no_report():
    """无重复行时不产生重复报告"""
    data_service.full_params_dup_report = {}
    df = _fp_df([
        {"eNodeB标识": "100", "小区标识": "1", "小区名称": "A"},
        {"eNodeB标识": "200", "小区标识": "2", "小区名称": "B", "基站经度": 114.0,
         "基站纬度": 24.0, "天线方向角": 60},
    ])
    data_service._parse_full_params_dataframe(df, "LTE")
    assert "LTE" not in data_service.full_params_dup_report
