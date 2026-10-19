"""工具类结果行的分页落盘与预览（站间距计算 / 批量距离匹配的页内列表）。

页内列表要「懒加载 + 内部滚动」，因此结果行必须能分页取，且不能整表塞进计算响应。
另外距离匹配的 Excel 与预览必须共用同一份行序——两处各写一份交错逻辑迟早会漂移。
"""
from __future__ import annotations

import pandas as pd
import pytest

from app.core.config import settings
from app.services import result_rows_store as store


@pytest.fixture()
def out_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "OUTPUT_DIR", tmp_path)
    return tmp_path


COLUMNS = ["站点名称", "经度", "纬度", "平均站间距(km)", "参与计算邻居数"]


def _rows(n: int):
    return [[f"cell-{i}", 113.0 + i * 1e-4, 23.0, round(i * 0.1, 4), 5] for i in range(n)]


def test_write_read_roundtrip(out_dir):
    store.write_rows("结果A", COLUMNS, _rows(5))

    columns, total, rows = store.read_page("结果A", offset=0, limit=200)
    assert columns == COLUMNS
    assert total == 5
    assert rows == _rows(5)


def test_pagination_windows(out_dir):
    store.write_rows("结果B", COLUMNS, _rows(10))

    _, total, page1 = store.read_page("结果B", offset=0, limit=4)
    assert total == 10 and [r[0] for r in page1] == ["cell-0", "cell-1", "cell-2", "cell-3"]

    _, _, page2 = store.read_page("结果B", offset=4, limit=4)
    assert [r[0] for r in page2] == ["cell-4", "cell-5", "cell-6", "cell-7"]

    # 末页不足一页
    _, _, page3 = store.read_page("结果B", offset=8, limit=4)
    assert [r[0] for r in page3] == ["cell-8", "cell-9"]


def test_offset_beyond_end_returns_empty(out_dir):
    store.write_rows("结果C", COLUMNS, _rows(3))
    _, total, rows = store.read_page("结果C", offset=99, limit=10)
    assert total == 3 and rows == []


def test_limit_is_clamped(out_dir):
    store.write_rows("结果D", COLUMNS, _rows(3))
    _, _, rows = store.read_page("结果D", offset=0, limit=10_000)
    assert len(rows) == 3  # 只有 3 行；上限约束由 MAX_PAGE_SIZE 保证
    assert store.MAX_PAGE_SIZE == 1000


def test_numpy_scalars_become_numbers(out_dir):
    """numpy 标量必须落成 JSON 数字：否则前端拿到的距离是字符串，无法排序。"""
    import numpy as np

    store.write_rows("结果E", COLUMNS, [[np.str_("cell-x"), np.float64(113.5), np.float64(23.5), np.float64(1.2345), np.int64(5)]])
    _, _, rows = store.read_page("结果E", offset=0, limit=10)
    assert isinstance(rows[0][3], float)
    assert isinstance(rows[0][4], int)
    assert rows[0][1] == 113.5


def test_missing_and_illegal_result_id(out_dir):
    store.write_rows("结果F", COLUMNS, _rows(1))

    with pytest.raises(store.ResultRowsError):
        store.read_page("不存在的结果")

    # 目录穿越必须被拒（即使目标文件真的存在）
    outside = out_dir.parent / "secret.rows.jsonl"
    outside.write_text('{"columns": [], "total": 0}\n', encoding="utf-8")
    for bad in ("../secret", "..\\secret", "sub/secret", "", "   "):
        with pytest.raises(store.ResultRowsError):
            store.read_page(bad)


def test_distance_match_preview_matches_excel_row_order(out_dir, tmp_path):
    """Excel 与预览必须同序同行：共用 build_output_rows。"""
    from openpyxl import load_workbook

    from app.services.distance_match_service import (
        OUTPUT_COLUMNS,
        build_output_excel,
        build_output_rows,
        match_sites,
    )

    match_df = pd.DataFrame({
        "CELL": ["m1", "m2", "m3"],
        "LON": [113.0, 113.5, 114.0],
        "LAT": [23.0, 23.0, 23.0],
    })
    full_df = pd.DataFrame({
        "CELL": ["f1", "f2", "f3"],
        "LON": [113.001, 113.502, 130.0],
        "LAT": [23.0, 23.0, 23.0],
    })
    results, unmatched = match_sites(match_df, full_df, max_distance_km=1.0)

    rows = build_output_rows(match_df, results, unmatched, 1.0)

    excel_path = tmp_path / "匹配结果.xlsx"
    build_output_excel(str(excel_path), match_df, results, unmatched, 1.0)
    ws = load_workbook(excel_path).active
    excel_rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]

    assert len(rows) == len(excel_rows) == 3  # m1/m2 各 1 条命中 + m3 备注行
    assert [r[0] for r in rows] == [r[0] for r in excel_rows] == ["m1", "m2", "m3"]
    # 无命中的站点保留为备注行
    assert rows[-1][7] and "无站点匹配" in rows[-1][7]

    store.write_rows("匹配结果", OUTPUT_COLUMNS, rows)
    columns, total, page = store.read_page("匹配结果", offset=0, limit=10)
    assert columns == OUTPUT_COLUMNS
    assert total == 3
    assert [r[0] for r in page] == [r[0] for r in excel_rows]
